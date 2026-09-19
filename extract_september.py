import glob
import openpyxl


path = glob.glob("/tmp/sales-read.*/*.xlsx")[0]
formula_book = openpyxl.load_workbook(path, data_only=False)
value_book = openpyxl.load_workbook(path, data_only=True)

for formula_sheet in formula_book.worksheets:
    value_sheet = value_book[formula_sheet.title]
    starts = []
    for row in formula_sheet.iter_rows():
        if any(isinstance(cell.value, str) and "กันยายน" in cell.value for cell in row):
            starts.append(row[0].row)
    if not starts:
        continue
    print(f"### {formula_sheet.title}")
    for start in starts:
        print(f"BLOCK {start}")
        for row_number in range(max(1, start - 2), min(start + 38, formula_sheet.max_row) + 1):
            parts = []
            for column_number in range(1, min(formula_sheet.max_column, 30) + 1):
                formula_cell = formula_sheet.cell(row_number, column_number)
                if formula_cell.value is None:
                    continue
                value = formula_cell.value
                if isinstance(value, str) and value.startswith("="):
                    value = value_sheet.cell(row_number, column_number).value
                parts.append(f"{formula_cell.coordinate}={value}")
            if parts:
                print(" | ".join(parts))
