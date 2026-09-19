import datetime as dt
import glob
import json
import re
import calendar
from collections import Counter, defaultdict

import openpyxl


SOURCE = glob.glob("/tmp/sales-read.*/*.xlsx")[0]
OUTPUT = "tmp/september_sales.json"
YEAR = 2026
MONTH = 9
PRODUCT_CODES = {
    "เล็ก": "01",
    "โม่": "02",
    "ก้อน": "05",
}
REGULAR_CODE = re.compile(r"^(?:AA|BB|CC)\d+$|^SW-\d+$|^B-ISO-\d+$", re.I)


def text(value):
    return "" if value is None else str(value).strip()


def normalized_code(value):
    return re.sub(r"\s+", "", text(value)).upper()


def numeric(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and value > 0:
        return float(value)
    return None


formula_book = openpyxl.load_workbook(SOURCE, data_only=False)
value_book = openpyxl.load_workbook(SOURCE, data_only=True)
records = []

for formula_sheet in formula_book.worksheets:
    value_sheet = value_book[formula_sheet.title]
    month_rows = []
    for row in formula_sheet.iter_rows():
        if any(isinstance(cell.value, str) and "กันยายน" in cell.value for cell in row):
            month_rows.append(row[0].row)

    for block_index, month_row in enumerate(month_rows):
        date_row = month_row + 1
        header_row = month_row + 2
        next_month_row = month_rows[block_index + 1] if block_index + 1 < len(month_rows) else formula_sheet.max_row + 1

        money_columns = [
            cell.column for cell in formula_sheet[header_row]
            if text(cell.value) == "เงิน"
        ]
        if not money_columns:
            continue

        groups = []
        previous_end = 2
        for money_column in money_columns:
            group_start = previous_end + 1
            product_labels = {
                text(formula_sheet.cell(header_row, column).value)
                for column in range(group_start, money_column)
            }
            date_candidates = []
            for column in range(group_start, money_column + 1):
                value = value_sheet.cell(date_row, column).value
                if isinstance(value, (int, float)) and 1 <= int(value) <= 31:
                    date_candidates.append(int(value))
            # Old weekly blocks sometimes retain hidden/formatted "เงิน" columns
            # to the left of the current block.  Only accept a group that has an
            # actual product header in this block.
            if date_candidates and product_labels.intersection(PRODUCT_CODES):
                # The first group can span stale columns from the prior printed
                # week. The date nearest the current money column is authoritative.
                day = date_candidates[-1]
                if day <= calendar.monthrange(YEAR, MONTH)[1]:
                    groups.append((group_start, money_column, day))
            previous_end = money_column

        if not groups:
            continue

        for row_number in range(header_row + 1, min(next_month_row, month_row + 80)):
            code = normalized_code(value_sheet.cell(row_number, 1).value)
            name = text(value_sheet.cell(row_number, 2).value)
            if code == "ลำดับ":
                break
            if any(
                isinstance(cell.value, str)
                and re.search(r"เดือน.*2569", cell.value)
                for cell in formula_sheet[row_number]
            ):
                break
            if not code and not name and any(
                isinstance(cell.value, str) and cell.value.startswith("=SUM(")
                for cell in formula_sheet[row_number]
            ):
                break
            if name == "รวม" or code == "รวม":
                break
            if not code and not name:
                continue

            lowered_name = name.casefold()
            if lowered_name == "สด":
                kind = "casual"
            elif code.startswith("EV"):
                kind = "event"
            elif REGULAR_CODE.match(code):
                kind = "regular"
            else:
                kind = "unmapped_named"

            for group_start, money_column, day in groups:
                items = defaultdict(float)
                ignored = {}
                for column in range(group_start, money_column):
                    label = text(formula_sheet.cell(header_row, column).value)
                    quantity = numeric(value_sheet.cell(row_number, column).value)
                    if quantity is None:
                        continue
                    if label in PRODUCT_CODES:
                        items[PRODUCT_CODES[label]] += quantity
                    else:
                        ignored[label or f"col_{column}"] = quantity

                if not items and not ignored:
                    continue

                money = value_sheet.cell(row_number, money_column).value
                records.append({
                    "sheet": formula_sheet.title,
                    "row": row_number,
                    "date": dt.date(YEAR, MONTH, day).isoformat(),
                    "kind": kind,
                    "code": code,
                    "name": name,
                    "items": dict(sorted(items.items())),
                    "ignored_quantities": ignored,
                    "money": money,
                })

summary = {
    "source": SOURCE,
    "record_count": len(records),
    "counts_by_kind": Counter(record["kind"] for record in records),
    "dates": sorted({record["date"] for record in records}),
    "regular_codes": sorted({record["code"] for record in records if record["kind"] == "regular"}),
    "event_names": sorted({record["name"] for record in records if record["kind"] == "event"}),
    "unmapped_names": sorted({record["name"] for record in records if record["kind"] == "unmapped_named"}),
    "records": records,
}

with open(OUTPUT, "w", encoding="utf-8") as handle:
    json.dump(summary, handle, ensure_ascii=False, indent=2, default=str)

print(json.dumps({key: value for key, value in summary.items() if key != "records"}, ensure_ascii=False, indent=2, default=str))
